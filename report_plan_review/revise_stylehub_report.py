"""Revise figure descriptions and add monochrome low-level architecture diagrams."""

from pathlib import Path
import re
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("D:/PROJECT/StyleHub_Project_Report.docx")
OUTPUT = ROOT / "report" / "StyleHub_Project_Report_Revised.docx"
ASSET_DIR = ROOT / "report" / "low_level_architecture"

FIGURES = [
    ("4.1", "StyleHub Class Diagram"),
    ("6.1", "StyleHub Entity Relationship Diagram"),
    ("7.1", "StyleHub System Architecture"),
    ("8.0.1", "Overall Request Flow of StyleHub"),
    ("8.1.1", "Overall Project Folder"), ("8.1.2", "Backend Folder"),
    ("8.1.3", "Backend Config Folder"), ("8.1.4", "Backend Routes Folder"),
    ("8.1.5", "Backend Controllers Folder"), ("8.1.6", "Backend Services Folder"),
    ("8.1.7", "Frontend Folder"), ("8.1.8", "Frontend Source Folder"),
    ("8.1.9", "Frontend Pages Folder"), ("8.1.10", "Frontend Components Folder"),
    ("8.1.11", "Frontend Context Folder"), ("8.1.12", "Frontend Services Folder"),
    ("8.2.1", "Registration Page"), ("8.2.2", "Login Page"),
    ("8.2.3", "Profile and Username Page"), ("8.2.4", "Admin Login Page"),
    ("8.3.1", "Home and Featured Products"), ("8.3.2", "Product Catalogue"),
    ("8.3.3", "Product Detail"), ("8.4.1", "Add to Cart Action"),
    ("8.4.2", "Cart Items and Quantities"), ("8.4.3", "Cart Reward Summary"),
    ("8.4.4", "Checkout Page"), ("8.5.3", "Purchase History"),
    ("8.5.4", "Notification Panel"), ("8.6.1", "Personal Wishlist"),
    ("8.6.2", "Friend Search and Requests"), ("8.6.3", "Friends List"),
    ("8.6.4", "Shared Wishlist"), ("8.6.5", "Share Feedback"),
    ("8.7.3", "Style Assistant Chat"), ("8.8.1", "Virtual Try On Entry"),
    ("8.8.2", "Photo and Garment Selection"), ("8.8.3", "Generated Try On Result"),
    ("8.8.4", "Try On History"), ("8.9.1", "Admin Dashboard Overview"),
    ("8.9.2", "Admin Product Management"), ("8.9.3", "Admin Users and Purchases"),
    ("8.9.4", "Admin Activity and VTON Metrics"), ("8.10.1", "Settings Drawer"),
    ("8.10.2", "Language Selector"), ("8.10.3", "Theme Selector"),
    ("8.10.4", "Communication Preferences"),
]

ARCHITECTURES = {
    "8.2.4 Low Level Architecture": (
        "8.2.4.1", "Authentication and Profile Low Level Architecture",
        ["Client React\nLogin and Profile Pages", "HTTP Axios and JWT\nProtected API Request",
         "Express Middleware\nRate Limit and Auth Check", "Auth Controller\nRegister Login Profile",
         "Credential Service\nbcrypt and JWT", "SQLite Users Table\nUser and Role Records"],
    ),
    "8.3.4 Low Level Architecture": (
        "8.3.4.1", "Product Catalogue Low Level Architecture",
        ["Client React\nCatalogue and Product Detail", "HTTP Axios\nProduct and Recent View Request",
         "Express Middleware\nJWT for Signed In Actions", "Product Controller\nList Detail Recent Views",
         "Product Service\nAvailability and View Logic", "SQLite Products and Recently Accessed\nCatalogue and View Records"],
    ),
    "8.4.4 Low Level Architecture": (
        "8.4.4.1", "Cart and Checkout Low Level Architecture",
        ["Client React\nProduct Cart and Checkout Pages", "HTTP Axios and JWT\nCart and Purchase Request",
         "Express Middleware\nAuthentication Check", "Cart and Purchase Controllers\nItems Quantities Checkout",
         "Purchase Service\nReward Total and Stock Validation", "SQLite Cart Purchase and Product Tables\nCart Purchase and Stock Records"],
    ),
    "8.5.4 Low Level Architecture": (
        "8.5.4.1", "Purchases and Notifications Low Level Architecture",
        ["Client React\nCheckout History Notifications", "HTTP Axios and JWT\nPurchase and Notification Request",
         "Express Middleware\nAuthentication Check", "Purchase Notification Controllers\nSave Read and Retrieve",
         "Communication Service\nIn App and Configured Channels", "SQLite Purchase Notification Delivery Tables\nTransaction and Delivery Records"],
    ),
    "8.6.4 Low Level Architecture": (
        "8.6.4.1", "Wishlist and Friends Low Level Architecture",
        ["Client React\nWishlist Friends and Share Pages", "HTTP Axios and JWT\nSocial Feature Request",
         "Express Middleware\nAuthentication Check", "Wishlist and Friends Controllers\nSave Request Share Feedback",
         "Social Service\nPermission and Status Logic", "SQLite Wishlist Friendship Share Tables\nSocial Relationship Records"],
    ),
    "8.7.4 Low Level Architecture": (
        "8.7.4.1", "Recommendations and Chatbot Low Level Architecture",
        ["Client React\nProduct Page and Style Assistant", "HTTP Axios\nRecommendation and Chat Request",
         "Express Middleware\nInput and Token Check", "Recommendation Chatbot Controllers\nRoute Requests",
         "Rule Engine\nMatching and Intent Logic", "SQLite Products and Recommendation Tables\nCatalogue and Rule Data"],
    ),
    "8.8.4 Low Level Architecture": (
        "8.8.4.1", "Virtual Try On Low Level Architecture",
        ["Client React\nTry On Page and Image Selection", "HTTP Multipart Request\nImage and Product Data",
         "Express Middleware\nUpload Validation and Auth", "VTON Controller\nProduct and Asset Validation",
         "Local Python FASHN VTON Service\nGenerate Output Image", "SQLite Try On Results Table\nSaved Output and Status"],
    ),
    "8.9.4 Low Level Architecture": (
        "8.9.4.1", "Admin Dashboard Low Level Architecture",
        ["Client React\nAdmin Dashboard and Management", "HTTP Axios and JWT\nRestricted Admin Request",
         "Express Middleware\nAuth and ADMIN Role Check", "Admin Controller\nMetrics Users Purchases Products",
         "Admin Service\nValidation and Summary Queries", "SQLite User Product Purchase Tables\nAdministrative Records"],
    ),
    "8.10.4 Low Level Architecture": (
        "8.10.4.1", "Theme Language and Preferences Low Level Architecture",
        ["Client React\nSettings Drawer and Controls", "Language Theme Context\nImmediate Interface Update",
         "HTTP Axios and JWT\nPreference Request", "Settings Controller\nSupported Values and User Scope",
         "Preference Service\nLocal Persistence and Validation", "SQLite Users and Notification Delivery Tables\nAccount Preference Records"],
    ),
}


def set_run_font(run, size=11, italic=False, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold


def set_character_spacing(run, value="10"):
    rpr = run._element.get_or_add_rPr()
    spacing = rpr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        rpr.append(spacing)
    spacing.set(qn("w:val"), value)  # 0.5 pt in twentieths of a point.


def clear_text(paragraph):
    for run in paragraph.runs:
        for text_node in run._element.findall(qn("w:t")):
            text_node.text = ""
        for br in list(run._element.findall(qn("w:br"))):
            run._element.remove(br)


def insert_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def add_caption(paragraph, text):
    clear_text(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=10, italic=True)


def add_description(paragraph, title):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    text = [
        "This figure shows the main screen or structure used by the module.",
        "It identifies the principal components visible in the workflow.",
        "These components process the relevant request and return a result.",
        "It links the implemented design to the documented module behaviour.",
    ]
    for index, sentence in enumerate(text):
        run = paragraph.add_run(sentence)
        set_run_font(run, size=11)
        set_character_spacing(run)
        if index < len(text) - 1:
            run.add_break()


def create_drawio(path, title, blocks):
    mxfile = ET.Element("mxfile", {
        "host": "app.diagrams.net", "modified": "2026-10-07T00:00:00.000Z",
        "agent": "StyleHub report", "version": "24.7.17", "type": "device",
    })
    diagram = ET.SubElement(mxfile, "diagram", {"id": path.stem, "name": title})
    model = ET.SubElement(diagram, "mxGraphModel", {
        "dx": "850", "dy": "1200", "grid": "1", "gridSize": "10", "page": "1",
        "pageScale": "1", "pageWidth": "827", "pageHeight": "1169",
        "background": "#FFFFFF", "math": "0", "shadow": "0",
    })
    root = ET.SubElement(model, "root")
    ET.SubElement(root, "mxCell", {"id": "0"})
    ET.SubElement(root, "mxCell", {"id": "1", "parent": "0"})
    for index, block in enumerate(blocks):
        ident = f"box-{index}"
        cell = ET.SubElement(root, "mxCell", {
            "id": ident, "value": block.replace("\n", "<br>"),
            "style": ("rounded=1;whiteSpace=wrap;html=1;fillColor=#FFFFFF;"
                      "strokeColor=#000000;fontColor=#000000;fontFamily=Times New Roman;"
                      "fontSize=13;fontStyle=1;align=center;verticalAlign=middle;"),
            "vertex": "1", "parent": "1",
        })
        ET.SubElement(cell, "mxGeometry", {
            "x": "100", "y": str(80 + index * 165), "width": "620", "height": "105",
            "as": "geometry",
        })
        if index:
            edge = ET.SubElement(root, "mxCell", {
                "id": f"edge-{index}", "style": ("endArrow=block;html=1;rounded=0;"
                                                  "strokeColor=#000000;"),
                "edge": "1", "parent": "1", "source": f"box-{index - 1}", "target": ident,
            })
            ET.SubElement(edge, "mxGeometry", {"relative": "1", "as": "geometry"})
    ET.ElementTree(mxfile).write(path, encoding="utf-8", xml_declaration=True)


def draw_diagram(path, blocks):
    image = Image.new("RGB", (1000, 1240), "white")
    draw = ImageDraw.Draw(image)
    font_file = Path("C:/Windows/Fonts/times.ttf")
    bold_file = Path("C:/Windows/Fonts/timesbd.ttf")
    font_file = Path(r"C:WindowsFonts	imes.ttf")
    bold_file = Path(r"C:WindowsFonts	imesbd.ttf")
    font_file = Path("C:/Windows/Fonts/times.ttf")
    bold_file = Path("C:/Windows/Fonts/timesbd.ttf")
    regular = ImageFont.truetype(str(font_file), 25)
    bold = ImageFont.truetype(str(bold_file if bold_file.exists() else font_file), 28)
    x1, x2 = 95, 905
    y = 65
    for index, block in enumerate(blocks):
        draw.rounded_rectangle((x1, y, x2, y + 115), radius=12, fill="white", outline="black", width=3)
        title, *body = block.split("\n")
        draw.text(((x1 + x2) / 2, y + 35), title, font=bold, fill="black", anchor="mm")
        if body:
            draw.text(((x1 + x2) / 2, y + 76), "\n".join(body), font=regular, fill="black", anchor="mm", align="center")
        if index < len(blocks) - 1:
            start, end = y + 115, y + 155
            draw.line((500, start, 500, end), fill="black", width=3)
            draw.polygon([(500, end + 12), (490, end - 3), (510, end - 3)], fill="black")
        y += 185
    image.save(path)


def replace_existing_figure_captions(document):
    lookup = {number: title for number, title in FIGURES}
    figure_pattern = re.compile(r"Fig\s+(\d+(?:\.\d+)+)")
    heading_pattern = re.compile(r"^(?:\d+\.\d+|CHAPTER\s+\d+|CONCLUSION|REFERENCES)")
    original_paragraphs = list(document.paragraphs)
    for position, paragraph in enumerate(original_paragraphs):
        matches = list(figure_pattern.finditer(paragraph.text))
        if not matches:
            continue
        numbers = []
        for match in matches:
            number = match.group(1)
            if number in lookup and number not in numbers:
                numbers.append(number)
        if not numbers:
            continue
        paragraph.paragraph_format.keep_with_next = True
        clear_text(paragraph)
        anchor = paragraph
        for number in numbers:
            caption_paragraph = insert_after(anchor)
            add_caption(caption_paragraph, f"Fig {number} {lookup[number]}")
            description_paragraph = insert_after(caption_paragraph)
            add_description(description_paragraph, lookup[number])
            anchor = description_paragraph
        # Remove any old prose that belonged to the figure. Image-only
        # paragraphs are retained so user-provided screenshots are preserved.
        for following in original_paragraphs[position + 1:]:
            text = following.text.strip()
            if figure_pattern.search(text) or heading_pattern.match(text):
                break
            if text and "w:drawing" not in following._p.xml:
                clear_text(following)


def insert_architecture_figures(document):
    for paragraph in list(document.paragraphs):
        heading = paragraph.text.strip()
        if heading not in ARCHITECTURES:
            continue
        number, title, blocks = ARCHITECTURES[heading]
        stem = number.replace(".", "_")
        drawio_path = ASSET_DIR / f"{stem}_low_level_architecture.drawio"
        png_path = ASSET_DIR / f"{stem}_low_level_architecture.png"
        create_drawio(drawio_path, title, blocks)
        draw_diagram(png_path, blocks)
        image_paragraph = insert_after(paragraph)
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(4)
        image_paragraph.paragraph_format.space_after = Pt(2)
        image_paragraph.paragraph_format.keep_with_next = True
        image_paragraph.add_run().add_picture(str(png_path), width=Cm(14.2))
        caption_paragraph = insert_after(image_paragraph)
        add_caption(caption_paragraph, f"Fig {number} {title}")
        description_paragraph = insert_after(caption_paragraph)
        add_description(description_paragraph, title)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    document = Document(SOURCE)
    replace_existing_figure_captions(document)
    insert_architecture_figures(document)
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("Fig "):
            paragraph.paragraph_format.keep_with_next = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
