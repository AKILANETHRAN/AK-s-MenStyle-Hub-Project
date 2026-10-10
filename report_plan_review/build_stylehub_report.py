from pathlib import Path
import re
import sqlite3
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from update_class_diagram import create_class_diagram
from update_er_diagram import create_er_diagram


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'report'
OUT.mkdir(exist_ok=True)
DB = ROOT / 'stylehub/stylehub-backend/config/stylehub.sqlite'


def diagram(path, boxes, edges, width, height):
    mxfile = ET.Element('mxfile', {'host': 'app.diagrams.net', 'modified': '2026-10-06T00:00:00.000Z', 'agent': 'StyleHub report', 'version': '24.7.17', 'type': 'device'})
    page = ET.SubElement(mxfile, 'diagram', {'id': 'stylehub', 'name': 'Page 1'})
    model = ET.SubElement(page, 'mxGraphModel', {'dx': '1200', 'dy': '800', 'grid': '1', 'gridSize': '10', 'guides': '1', 'tooltips': '1', 'connect': '1', 'arrows': '1', 'fold': '1', 'page': '1', 'pageScale': '1', 'pageWidth': str(width), 'pageHeight': str(height), 'math': '0', 'shadow': '0'})
    root = ET.SubElement(model, 'root')
    ET.SubElement(root, 'mxCell', {'id': '0'})
    ET.SubElement(root, 'mxCell', {'id': '1', 'parent': '0'})
    for key, label, x, y, w, h in boxes:
        cell = ET.SubElement(root, 'mxCell', {'id': key, 'value': label.replace('\n', '<br>'), 'style': 'rounded=0;whiteSpace=wrap;html=1;fillColor=#FFFFFF;strokeColor=#000000;fontColor=#000000;fontFamily=Times New Roman;fontSize=15;align=center;verticalAlign=middle;spacing=7;', 'vertex': '1', 'parent': '1'})
        ET.SubElement(cell, 'mxGeometry', {'x': str(x), 'y': str(y), 'width': str(w), 'height': str(h), 'as': 'geometry'})
    for idx, (a, b, label) in enumerate(edges, 1):
        cell = ET.SubElement(root, 'mxCell', {'id': f'e{idx}', 'value': label, 'style': 'endArrow=none;html=1;rounded=0;strokeColor=#000000;fontColor=#000000;fontFamily=Times New Roman;fontSize=12;labelBackgroundColor=#FFFFFF;', 'edge': '1', 'parent': '1', 'source': a, 'target': b})
        ET.SubElement(cell, 'mxGeometry', {'relative': '1', 'as': 'geometry'})
    ET.ElementTree(mxfile).write(path, encoding='utf-8', xml_declaration=True)


create_class_diagram(OUT/'StyleHub_Class_Diagram.drawio')

er_boxes = [
    ('users','users\nPK id\nusername, email, role',40,40,250,100),
    ('products','products\nPK id\nname, category, price, stock',430,40,250,100),
    ('cart','cart\nPK id, FK user_id',820,40,250,85),
    ('cart_items','cart_items\nPK id, FK cart_id, product_id',820,200,250,85),
    ('wishlists','wishlists\nPK id, FK user_id',40,230,250,85),
    ('wishlist_items','wishlist_items\nPK id, FK wishlist_id, product_id',430,230,250,85),
    ('friendships','friendships\nPK id, FK requester_id, receiver_id',40,420,250,95),
    ('shares','wishlist_shares\nPK id, FK sender_id, receiver_id',430,420,250,90),
    ('feedback','wishlist_feedback\nPK id, FK share_id, user_id',820,420,250,90),
    ('purchases','purchases\nPK id, FK user_id\npurchase_type, total_amount',40,640,250,110),
    ('purchase_items','purchase_items\nPK id, FK purchase_id, product_id',430,640,250,90),
    ('notifications','notifications\nPK id, FK user_id',820,640,250,85),
    ('tryon','try_on_results\nPK id, FK user_id, product_id',40,850,250,90),
    ('recommend','outfit_recommendations\nPK id, FK base_product_id, recommended_product_id',430,850,250,110),
    ('recent','recently_accessed\nPK id, FK user_id, product_id',820,850,250,90),
]
er_edges = [('users','cart','1:1'),('cart','cart_items','1:N'),('products','cart_items','1:N'),('users','wishlists','1:1'),('wishlists','wishlist_items','1:N'),('products','wishlist_items','1:N'),('users','friendships','1:N'),('users','shares','1:N'),('shares','feedback','1:N'),('users','purchases','1:N'),('purchases','purchase_items','1:N'),('products','purchase_items','1:N'),('users','notifications','1:N'),('users','tryon','1:N'),('products','tryon','1:N'),('products','recommend','1:N'),('users','recent','1:N'),('products','recent','1:N')]
create_er_diagram(OUT/'StyleHub_ER_Diagram.drawio')

arch_boxes = [
    ('browser','User Browser\nReact and Vite Frontend',480,30,280,80),
    ('api','Express REST API\nJWT and route middleware',480,175,280,80),
    ('db','SQLite Database\nUsers, Products, Purchases',100,340,260,90),
    ('logic','Business Services\nCart and Recommendations\nChatbot',480,340,280,90),
    ('vton','Local Python VTON Service\nFASHN VTON 1.5',860,340,280,90),
    ('external','Optional External Services\nGoogle OAuth and Email\nSMS and WhatsApp',480,525,280,100),
]
arch_edges = [('browser','api','HTTP'),('api','db','read / write'),('api','logic','dispatch'),('api','vton','image request'),('logic','external','configured integrations')]
diagram(OUT/'StyleHub_Architecture.drawio', arch_boxes, arch_edges, 1240, 690)

# PNG rendering of the same architecture layout used in the editable Draw.io source.
scale=2
img=Image.new('RGB',(1240*scale,690*scale),'white')
dr=ImageDraw.Draw(img)
fontpath=Path('C:/Windows/Fonts/times.ttf')
font=ImageFont.truetype(str(fontpath),42) if fontpath.exists() else ImageFont.load_default()
small=ImageFont.truetype(str(fontpath),30) if fontpath.exists() else ImageFont.load_default()
positions={k:(x,y,w,h) for k,_,x,y,w,h in arch_boxes}
def center(k):
    x,y,w,h=positions[k]; return ((x+w/2)*scale,(y+h/2)*scale)
for a,b,label in arch_edges:
    x1,y1=center(a); x2,y2=center(b)
    dr.line((x1,y1,x2,y2),fill='black',width=3)
    cx,cy=(x1+x2)/2,(y1+y2)/2
    bb=dr.textbbox((0,0),label,font=small)
    dr.rectangle((cx-(bb[2]-bb[0])/2-6,cy-13,cx+(bb[2]-bb[0])/2+6,cy+14),fill='white')
    dr.text((cx,cy),label,font=small,fill='black',anchor='mm')
for _,label,x,y,w,h in arch_boxes:
    dr.rectangle((x*scale,y*scale,(x+w)*scale,(y+h)*scale),outline='black',fill='white',width=3)
    lines=label.split('\n')
    for i,line in enumerate(lines):
        dr.text(((x+w/2)*scale,(y+h/2)*scale+(i-(len(lines)-1)/2)*50),line,font=font,fill='black',anchor='mm')
img.save(OUT/'StyleHub_Architecture.png',dpi=(220,220))


def cell_shading(cell):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'FFFFFF'); tcPr.append(shd)

def set_spacing(run):
    rpr=run._element.get_or_add_rPr(); item=OxmlElement('w:spacing'); item.set(qn('w:val'),'10'); rpr.append(item)

doc=Document()
sec=doc.sections[0]; sec.page_height=Cm(29.7); sec.page_width=Cm(21); sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.2); sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)
styles=doc.styles
for name,size,bold,after in [('Normal',12,False,6),('Title',18,True,12),('Heading 1',16,True,12),('Heading 2',14,True,8),('Heading 3',12,True,6)]:
    s=styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(size); s.font.bold=bold; s.font.color.rgb=RGBColor(0,0,0)
    rfonts=s._element.get_or_add_rPr().rFonts
    if rfonts is not None:
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme'):
            rfonts.attrib.pop(qn(f'w:{attr}'),None)
        rfonts.set(qn('w:ascii'),'Times New Roman'); rfonts.set(qn('w:hAnsi'),'Times New Roman'); rfonts.set(qn('w:eastAsia'),'Times New Roman')
    s.paragraph_format.space_after=Pt(after); s.paragraph_format.line_spacing=1.3
styles['Normal'].paragraph_format.line_spacing=1.5
title_ppr=styles['Title']._element.get_or_add_pPr()
for old in title_ppr.findall(qn('w:pBdr')):
    title_ppr.remove(old)

def p(text='',style=None,align=None):
    x=doc.add_paragraph(style=style); x.add_run(text)
    if align is not None: x.alignment=align
    return x
def chapter(num,title):
    chapter_label=p(f'CHAPTER {num}',align=WD_ALIGN_PARAGRAPH.CENTER)
    chapter_label.paragraph_format.page_break_before=True
    chapter_label.paragraph_format.keep_with_next=True
    doc.add_heading(title,1)
def sub(title): doc.add_heading(title,2)
def body(text): p(text)
def bullet(text):
    item=p('•  '+text)
    item.paragraph_format.left_indent=Cm(0.7)
    item.paragraph_format.first_line_indent=Cm(-0.35)
    return item
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): cells[i].text=str(val)
    for row in t.rows:
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_shading(c)
            for cp in c.paragraphs:
                cp.paragraph_format.space_after=Pt(2); cp.paragraph_format.line_spacing=1.1
                for run in cp.runs: run.font.name='Times New Roman'; run.font.size=Pt(9)
    for c in t.rows[0].cells:
        for run in c.paragraphs[0].runs: run.bold=True
    header_pr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); header_pr.append(repeat)
    if widths:
        for row in t.rows:
            for i,c in enumerate(row.cells): c.width=Cm(widths[i])
    doc.add_paragraph()
    return t
def caption(text):
    cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after=Pt(3)
    r=cp.add_run(text); r.italic=True; r.font.name='Times New Roman'; r.font.size=Pt(10)
def desc(text):
    dp=doc.add_paragraph(); dp.paragraph_format.line_spacing=1.5; dp.paragraph_format.space_after=Pt(10)
    for i,line in enumerate(text.split('\n')):
        r=dp.add_run(line); r.font.name='Times New Roman'; r.font.size=Pt(11); set_spacing(r)
        if i < len(text.split('\n'))-1: r.add_break()
def screenshot(number,title,description):
    gap=doc.add_paragraph(' ')
    gap.paragraph_format.space_before=Pt(2)
    gap.paragraph_format.space_after=Cm(2.2)
    gap.paragraph_format.keep_with_next=True
    caption(f'Fig {number} {title}')
    desc(description)

# Editable preliminary pages. The student fills the blank fields later.
p('STYLEHUB MEN FASHION E COMMERCE SYSTEM','Title',WD_ALIGN_PARAGRAPH.CENTER)
p('PROJECT REPORT',align=WD_ALIGN_PARAGRAPH.CENTER)
p('Submitted by',align=WD_ALIGN_PARAGRAPH.CENTER)
p('Name: ____________________    Register Number: ____________________',align=WD_ALIGN_PARAGRAPH.CENTER)
p('Course Code and Title: __________________________________________',align=WD_ALIGN_PARAGRAPH.CENTER)
p('Guide: ____________________',align=WD_ALIGN_PARAGRAPH.CENTER)
p('Department: ____________________',align=WD_ALIGN_PARAGRAPH.CENTER)
p('College: ____________________',align=WD_ALIGN_PARAGRAPH.CENTER)
p('Month and Year: ____________________',align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()
p('BONAFIDE CERTIFICATE','Title',WD_ALIGN_PARAGRAPH.CENTER)
body('This is to certify that the project report entitled "StyleHub Men Fashion E Commerce System" is the bonafide work of ____________________ (Register Number ____________________) carried out under my supervision.')
p('Guide Signature: ____________________')
p('Head of Department Signature: ____________________')
p('Date: ____________________')
doc.add_page_break()
p('TABLE OF CONTENTS','Title',WD_ALIGN_PARAGRAPH.CENTER)
table(['Chapter','Title'],[(i,t) for i,t in enumerate(['Introduction','Objectives','Description','Class Diagram','Table Structures','ER Diagram','Tech Stack and Architecture','Modules','Conclusion'],1)],[3,13])
body('Page numbers can be updated after the final screenshots and cover details are inserted.')

chapter(1,'INTRODUCTION')
body('Online fashion shopping gives customers access to many products, but choosing a suitable garment from pictures and product descriptions can still be difficult. A shopper must compare price, size, appearance, availability, and matching pieces while moving between multiple pages. A conventional catalogue provides the product record but does little to connect the selection, purchase, and personal styling decisions. StyleHub addresses this need through one men fashion web application that brings those activities together.')
body('The system allows a registered user to browse a catalogue, inspect product details, view recently accessed items, save favourites, manage a cart, and purchase individual products, a coordinated outfit, or cart contents. It also adds social features: users can find friends, share selected wishlist products privately, and exchange reactions and comments. A rule based recommendation component suggests complementary garments for eligible products, while a deterministic chatbot answers common product and shopping questions. These two components use defined application logic rather than a trained conversational model.')
body('For supported garments, StyleHub can send a customer photograph and a product garment image to a separate local FASHN VTON service and display the resulting virtual try-on image. This feature depends on the Python service, its model weights, and available hardware. The core shopping application continues to provide catalogue and purchase functions even when try-on is unavailable. Settings cover language, appearance, and notification preferences. An administrator has a restricted dashboard for product updates and operational summaries.')
body('The application is implemented with a React and Vite frontend, a Node.js and Express REST API, and an SQLite database. JWT authentication identifies the user on protected requests, while the database records products, carts, purchases, social relationships, notifications, and try-on results. Optional provider integrations support Google sign-in and purchase communication when configured. This report documents the implemented modules, their data structures and APIs, and the interfaces to be illustrated by screenshots during final submission.')

chapter(2,'OBJECTIVES')
for x in ['To provide secure registration, login, profile management, and protected user and administrator pages.','To offer a men fashion catalogue with product details, stock, prices, and recently accessed items.','To support size aware cart management, applicable rewards, checkout, and purchase history.','To let users save favourites and revisit products without repeating a catalogue search.','To allow users to connect with friends and privately share selected wishlist products.','To record reactions and comments on shared selections for collaborative style decisions.','To generate rule based complementary outfit suggestions for eligible products.','To answer common product and shopping questions through a deterministic assistant.','To offer virtual try-on for supported garments through a local Python service when available.','To maintain notification records and user communication preferences.','To give administrators a restricted dashboard for metrics, users, purchases, and product updates.','To support English, Tamil, and Hindi interface options with a saved language choice.','To keep user purchases and social records separated by authenticated account.']:
    bullet(x)

chapter(3,'DESCRIPTION')
body('StyleHub is a full stack men fashion application consisting of the following major functional areas. Each area is connected through the same user account and product catalogue.')
sub('1 User Authentication and Access Control')
body('Registration and sign-in establish an account identity. Passwords are checked against stored hashes, and a JWT protects customer API requests. Profile editing and username completion support account maintenance, including Google-created accounts when that integration is configured. An ADMIN role controls access to the management dashboard.')
sub('2 Product Discovery and Personal Shopping')
body('The catalogue displays product names, descriptions, images, categories, prices, stock, and applicable virtual try-on support. Product detail pages expose purchase, cart, wishlist, and outfit suggestion actions. Recently accessed records help a signed-in shopper return to previously viewed products. The cart retains selected products, sizes, and quantities before checkout.')
sub('3 Purchase and Communication')
body('The purchase functions handle a single product, a coordinated combination, or all eligible cart items. They create purchase and purchase item records, calculate totals and applicable cart rewards, and update stock. A purchase history page lets the customer review saved transactions. In-app notifications and channel delivery records support purchase communication; external email, SMS, or WhatsApp delivery depends on configured providers.')
sub('4 Social Wishlist and Style Assistance')
body('Wishlists store favourites separately from a cart. A user can send a friend request, accept a connection, share selected products, and respond to a shared list with feedback. Rule based outfit suggestions pair eligible products with complementary items. The chatbot answers defined styling and catalogue questions using application rules and product data.')
sub('5 Virtual Try-On and Administration')
body('Virtual try-on accepts an uploaded customer image for a supported garment and requests local FASHN VTON inference. The generated result is saved for later viewing when the service succeeds. The administrator dashboard provides restricted access to metrics, catalogue updates, user summaries, purchase records, and try-on activity.')
sub('6 Technology and Preferences')
body('React and Vite render the browser interface, Express exposes REST endpoints, and SQLite persists the application data. The interface supports English, Tamil, and Hindi and keeps the selected language. Appearance and communication settings are stored through the user settings flow. The optional Google sign-in and messaging integrations require separate provider credentials.')

chapter(4,'CLASS DIAGRAM')
body('The accompanying editable Draw.io class diagram shows the main domain objects and their associations. User connects to cart, wishlist, purchase, friendship, and try-on records. Products are referenced by cart items, wishlist items, purchase items, and try-on results. The administrator is represented by the ADMIN role on a user account.')
caption('Fig 4.1 StyleHub Class Diagram')
desc('Open StyleHub_Class_Diagram.drawio to view and edit the UML structure. The diagram uses only black lines and text.')

chapter(5,'TABLE STRUCTURES')
body('The following entries are taken from the SQLite schema in the project. PK denotes primary key and FK denotes foreign key. The tables store both shopping records and optional feature records such as social sharing and virtual try-on.')
con=sqlite3.connect(f'file:{DB.as_posix()}?mode=ro',uri=True)
names=[row[0] for row in con.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
for name in names:
    info=list(con.execute(f'PRAGMA table_info("{name}")'))
    fks={r[3]:r[2] for r in con.execute(f'PRAGMA foreign_key_list("{name}")')}
    doc.add_heading(name.replace('_',' ').title(),2)
    rows=[]
    for _,field,dtype,notnull,default,pk in info:
        key='PK' if pk else ('FK → '+fks[field] if field in fks else '')
        rows.append((field,dtype or 'TEXT',key,'Yes' if notnull else 'No'))
    table(['Field','Type','Key','Required'],rows,[5,3,4,3])
con.close()

chapter(6,'ER DIAGRAM')
body('The accompanying editable Draw.io ER diagram summarizes the main foreign key relationships. It includes shopping, social, recommendation, notification, recently accessed, and try-on data. One user can have many purchases; each purchase has many purchase items. Product references link items and feature records back to the catalogue.')
caption('Fig 6.1 StyleHub Entity Relationship Diagram')
desc('Open StyleHub_ER_Diagram.drawio to inspect and edit the black-and-white entity relationship diagram.')

chapter(7,'TECH STACK AND ARCHITECTURE')
sub('Technology Stack')
table(['Layer','Technology','Purpose'],[
    ('Frontend','React 18, Vite 6, React Router','Screens, navigation, user state'),
    ('Backend','Node.js, Express 4','REST API and business rules'),
    ('Database','SQLite, better-sqlite3','Persistent records and transactions'),
    ('Authentication','JWT, bcryptjs','Tokens and password hashing'),
    ('Virtual try-on','Python Flask, FASHN VTON 1.5','Local image inference for supported garments'),
    ('Communication','Nodemailer, Twilio SMS, Meta WhatsApp API','Optional configured purchase notifications'),
],[3,5,8])
sub('Architecture Overview')
body('The React browser application communicates with Express through REST endpoints. Express applies authentication and role checks, reads and writes SQLite records, and invokes business services. It calls the local VTON process for supported garment try-on. Configured external services support Google sign-in and notifications.')
doc.add_picture(str(OUT/'StyleHub_Architecture.png'),width=Cm(16))
caption('Fig 7.1 StyleHub System Architecture')
desc('The diagram shows request flow between the browser, Express API, SQLite database, business services, local VTON process, and optional external integrations. The editable source is StyleHub_Architecture.drawio.')

chapter(8,'MODULES')
sub('8.0 Design Approach')
body('The modules are described in a consistent order so that their responsibilities and implementation can be compared. Each section starts with its purpose, then lists observable functions, quality requirements, technologies, request flow, API endpoints, database design, and interfaces. The descriptions are based on the project routes, screens, and SQLite schema.')
table(['Design concern','Approach used in StyleHub'],[
    ('Identity','JWT identifies the user on protected requests; ADMIN routes require an administrator role.'),
    ('Data integrity','SQLite foreign keys connect user, product, cart, purchase, social, and try-on records.'),
    ('Modularity','React pages call frontend services; Express routes delegate application operations.'),
    ('User isolation','The backend uses the authenticated user ID for account-specific records.'),
    ('Optional services','VTON and provider communication report availability or delivery outcome separately.'),
],[4,12])
doc.add_heading('Overall Request Flow',3)
body('A user action begins in a React page. The frontend sends an HTTP request through its API service. Express selects a route, checks authentication where required, then performs validation and database or service work. The API returns a status and data to the page, which updates the displayed state. The local VTON service and external providers are called only for features that need them.')
table(['Step','Component','Result'],[
    ('1','React page or component','Collects input and starts an action'),
    ('2','Frontend API service','Sends request and JWT when protected'),
    ('3','Express route and middleware','Selects handler and checks access'),
    ('4','Controller and service','Validates and applies business rules'),
    ('5','SQLite or optional service','Reads, writes, or generates a result'),
    ('6','API response and React state','Shows data, confirmation, or an error'),
],[2,6,8])
caption('Fig 8.0.1 Overall Request Flow of StyleHub')
desc('The table follows an action from the browser to the API.\nMiddleware verifies protected access before a handler runs.\nThe handler uses SQLite or an optional service when needed.\nThe response updates the user interface or reports a failure.')

sub('8.1 Folder Structure')
body('stylehub/stylehub-frontend contains React pages, reusable components, contexts, and frontend service calls. stylehub/stylehub-backend contains Express routes, controllers, services, middleware, database configuration, and tests. The VTON service is housed under the backend vton-service directory.')
folder_figures=[
('8.1.1','Overall Project Folder','The top-level project separates the application and report files.\nThe frontend and backend sit under the stylehub directory.\nEach side has its own package and source organization.\nThis view helps locate the major project areas.'),
('8.1.2','Backend Folder','The backend contains server.js and its application packages.\nRoutes receive requests and controllers perform operations.\nServices hold shared rules and integrations.\nThe config directory contains SQLite initialization.'),
('8.1.3','Backend Config Folder','The config folder holds database and JWT support files.\nDatabase initialization defines the tables and indexes.\nThe SQLite file persists local application records.\nConfiguration is kept separate from request handlers.'),
('8.1.4','Backend Routes Folder','Route files group endpoints by feature.\nEach route maps an HTTP method and path to a handler.\nProtected routes apply authentication middleware.\nThe groups include users, shopping, social functions and administration.'),
('8.1.5','Backend Controllers Folder','Controllers validate inputs and assemble responses.\nEach controller belongs to a feature area.\nThey use database and service functions as needed.\nThis separation keeps route definitions concise.'),
('8.1.6','Backend Services Folder','Services contain reusable business and integration logic.\nRecommendation rules and communication handling are examples.\nThey support the controllers without defining page layout.\nProvider behaviour depends on configuration.'),
('8.1.7','Frontend Folder','The frontend package contains the Vite application.\nPublic assets support product and interface imagery.\nThe src directory contains executable React code.\nIts package file identifies React and build dependencies.'),
('8.1.8','Frontend Source Folder','The src directory separates pages, components and contexts.\nServices centralize browser requests to the API.\nThe application route map is defined in App.jsx.\nThis layout makes feature screens easier to find.'),
('8.1.9','Frontend Pages Folder','Pages represent the main browser routes.\nShopping, social, profile, admin and try-on screens appear here.\nProtectedRoute guards signed-in and ADMIN pages.\nA page can reuse components and contexts.'),
('8.1.10','Frontend Components Folder','Components hold reusable interface elements.\nExamples include the layout, product card and chatbot.\nShared display logic stays out of page files.\nComponents receive data from props and contexts.'),
('8.1.11','Frontend Context Folder','Contexts hold state shared across screens.\nAuthentication, cart, wishlist and friends use contexts.\nLanguage and appearance preferences are also managed here.\nThis avoids passing state through many page levels.'),
('8.1.12','Frontend Services Folder','Service modules call the backend API.\nThey group cart, wishlist, friends, notification and try-on requests.\nPages use these functions to perform actions.\nResponse data is then reflected in the interface.'),
]
for number,title,description in folder_figures:
    screenshot(number,title,description)

modules=[
{
 'n':'8.2','title':'Authentication and Profile','scope':'This module creates and maintains customer identities, protects account data and distinguishes an administrator from a customer. It covers registration, credential login, profile updates, username availability and the optional Google sign-in path.',
 'fr':['Register a user with required identity and credential details.','Authenticate a returning user and return a signed session token.','Load and update only the signed-in user profile.','Allow a Google-created account to complete a unique username.','Restrict administrator pages and APIs to accounts with the ADMIN role.'],
 'nfr':['Store password hashes rather than plaintext passwords.','Reject missing or invalid tokens on protected endpoints.','Keep one user from editing another user profile.'],
 'stack':'React, React Router, Express, JWT, bcryptjs and SQLite.',
 'flow':['The registration or login page sends form data to an auth route.','The controller validates input, checks the users table and creates or verifies credentials.','A successful login returns a JWT; ProtectedRoute and API middleware use it on later requests.','An ADMIN role is checked again on the backend before dashboard operations.'],
 'api':[('POST','/api/auth/register','Create customer account'),('POST','/api/auth/login','Authenticate credentials'),('GET','/api/auth/me','Return current user'),('PUT','/api/auth/profile','Update current profile'),('GET','/api/auth/check-username','Check username availability'),('PUT','/api/auth/username','Save username')],
 'db':[('users','id; email; username; password_hash; role','Account and access attributes')],
 'screens':[('Registration Page','Shows account fields and validation feedback.','Collects a new customer identity.','Submits the details to the registration API.','Successful registration enables sign-in.'),('Login Page','Shows the credential entry form.','Submits the account email and password.','Displays a response if authentication fails.','Successful login opens protected shopping pages.'),('Profile and Username Page','Displays the current account details.','Allows supported profile fields to be changed.','Checks a proposed username for availability.','Stores the completed username for the account.'),('Admin Login Page','Provides a separate administrator entry point.','Uses the existing account role for access.','Rejects a user without ADMIN permission.','Leads to the restricted dashboard after login.')]
},
{
 'n':'8.3','title':'Product Catalogue and Recent Views','scope':'This module presents products for discovery and records which items a signed-in user has viewed. Product records include catalogue text, images, prices, available stock and virtual try-on eligibility.',
 'fr':['List available catalogue products for browsing.','Open a detail page for a selected product.','Show category, price, discount, stock and garment imagery.','Record and retrieve recently accessed items per user.','Expose the selected product to cart, wishlist, recommendation and try-on actions.'],
 'nfr':['Return a clear result for an invalid product identifier.','Keep recently accessed records isolated by user.','Show stock and availability from persisted product data.'],
 'stack':'React pages and product components, Express product routes, SQLite.',
 'flow':['The catalogue page requests the product list from the Express API.','Selecting an item requests its detailed record.','The recent-access route updates the signed-in user and product pair.','The page renders available actions based on the returned product fields.'],
 'api':[('GET','/api/products','List products'),('GET','/api/products/:id','Get product detail'),('GET','/api/recently-accessed','List user recent views'),('POST','/api/recently-accessed','Record a product view')],
 'db':[('products','id; name; category; price; stock; image; vton_supported','Catalogue source'),('recently_accessed','user_id; product_id; accessed_at','Per-user viewing history')],
 'screens':[('Home and Featured Products','Introduces the fashion collection.','Shows selected catalogue entries.','Provides navigation to the full shop.','Shows recent items when available.'),('Product Catalogue','Presents product cards for browsing.','Shows image, name, price and stock state.','Allows opening an individual item.','Supports the discovery stage of shopping.'),('Product Detail','Shows details for a chosen product.','Presents garment attributes and actions.','Offers cart, wishlist or purchase choices.','May show recommendation or try-on availability.'),('Recently Accessed Products','Lists products viewed by this account.','Uses saved access timestamps for ordering.','Allows a shopper to reopen an item.','Keeps the history separate from other users.')]
},
{
 'n':'8.4','title':'Cart and Checkout','scope':'The cart groups selected products and size variants before purchase. Checkout calculates an order summary and applies the implemented Buy More Save More reward when the cart meets its rule conditions.',
 'fr':['Add a product and available size to the active cart.','Change quantity or remove an item.','Show subtotal and applicable cart reward.','Submit eligible cart contents with delivery details.','Create purchase records and decrement stock through the purchase path.'],
 'nfr':['Do not allow nonpositive quantities or unavailable stock to be purchased.','Keep each cart tied to its authenticated owner.','Perform stock and purchase updates in the same transaction where implemented.'],
 'stack':'React CartContext and checkout pages, Express cart and purchase routes, SQLite.',
 'flow':['The product page adds an item to the authenticated user cart.','Cart routes return current items and accept quantity or removal changes.','Checkout presents delivery details and computed totals.','The cart purchase route validates items, records the purchase and updates stock.'],
 'api':[('GET','/api/cart','Read current cart'),('POST','/api/cart/items','Add a cart item'),('PUT','/api/cart/items/:productId','Change quantity'),('DELETE','/api/cart/items/:productId','Remove item'),('POST','/api/purchases/cart','Purchase cart contents')],
 'db':[('cart','id; user_id','One active user cart'),('cart_items','cart_id; product_id; size; quantity','Selected variants'),('purchases','user_id; total_amount; cart_reward_amount','Completed purchase'),('purchase_items','purchase_id; product_id; size; quantity','Purchased units')],
 'screens':[('Add to Cart Action','Starts from a product detail view.','Collects a supported size if required.','Adds the selected product to the cart.','Updates the cart shown to the signed-in user.'),('Cart Items and Quantities','Lists products in the active cart.','Displays size and quantity for each line.','Supports quantity changes and removal.','Updates the total after a change.'),('Cart Reward Summary','Shows the current subtotal.','Displays the applicable reward calculation.','Makes the final amount visible before checkout.','Changes when eligible cart contents change.'),('Checkout Page','Collects the delivery details.','Shows the purchase summary.','Submits an authenticated cart purchase.','Returns a confirmation or validation message.')]
},
{
 'n':'8.5','title':'Purchases and Notifications','scope':'This module stores completed purchases and their items, provides a purchase history and records in-app and configured channel communication. Single, combination and cart purchase routes share the saved purchase model.',
 'fr':['Purchase one product, an eligible combination or a cart.','Save delivery data, purchased items and final amount.','Retrieve a signed-in user purchase history and one purchase detail.','Create and mark in-app notifications as read.','Record email, SMS or WhatsApp delivery attempts and outcomes.'],
 'nfr':['Return only purchases owned by the authenticated user.','Preserve the purchase when a nonessential external notification fails.','Record configured, disabled or failed delivery states accurately.'],
 'stack':'Express purchase and notification routes, SQLite, Nodemailer, Twilio SMS and Meta WhatsApp when configured.',
 'flow':['A purchase route validates product selection, stock and delivery fields.','The database stores a purchase and its line items, with associated stock changes.','The notification service creates an in-app message and attempts enabled channels.','History routes return records for the authenticated account.'],
 'api':[('POST','/api/purchases/product','Purchase one product'),('POST','/api/purchases/combo','Purchase an outfit'),('POST','/api/purchases/cart','Purchase cart'),('GET','/api/purchases','List own purchases'),('GET','/api/purchases/:id','Get one purchase'),('GET','/api/notifications','List notifications'),('PUT','/api/notifications/:id/read','Mark notification read')],
 'db':[('purchases','id; user_id; purchase_type; total_amount','Purchase header'),('purchase_items','purchase_id; product_id; quantity; unit_price','Purchase lines'),('notifications','user_id; title; message; is_read','In-app messages'),('notification_deliveries','purchase_id; channel; status; provider_message_id','Channel delivery audit')],
 'screens':[('Single Product Purchase','Starts from one selected product.','Collects delivery details for checkout.','Shows the price before submission.','Creates a purchase for that item.'),('Complete Look Purchase','Shows the coordinated garment selection.','Allows purchasing the eligible combination.','Calculates the combined amount.','Stores each selected product as a purchase item.'),('Purchase History','Lists saved purchases for the user.','Shows amount and purchase information.','Allows opening a selected record.','Keeps records tied to the signed-in account.'),('Notification Panel','Lists in-app purchase messages.','Shows read and unread state.','Allows a message to be marked read.','External channel delivery is audited separately.')]
},
{
 'n':'8.6','title':'Wishlist and Friends','scope':'The social module stores favourite products, manages friendship requests and supports private sharing with feedback. It keeps a user wishlist separate from the cart and limits sharing to permitted relationships.',
 'fr':['Add and remove products from the personal wishlist.','Find users and send, accept or reject friend requests.','Share selected wishlist items with a friend.','Open received shares and respond with reactions or comments.','Show the current friend and request lists.'],
 'nfr':['Restrict wishlist and share access to the relevant user.','Prevent duplicate wishlist items and invalid self-friendship records.','Show a clear status for pending, accepted or rejected requests.'],
 'stack':'React wishlist and friends pages, Express routes and controllers, SQLite.',
 'flow':['A user saves a product through the wishlist route.','The friends route stores a request and later updates its status.','An accepted connection can receive a selected product share.','Feedback is stored against the share and returned to its participants.'],
 'api':[('GET','/api/wishlist','Read wishlist'),('POST','/api/wishlist/items','Save product'),('DELETE','/api/wishlist/items/:productId','Remove product'),('POST','/api/friends/request','Send request'),('POST','/api/friends/:id/accept','Accept request'),('POST','/api/wishlist/shares','Create share'),('POST','/api/wishlist/shares/:shareId/feedback','Add feedback')],
 'db':[('wishlists; wishlist_items','user_id; wishlist_id; product_id','Saved products'),('friendships','requester_id; receiver_id; status','Relationship state'),('wishlist_shares; wishlist_share_items','sender_id; receiver_id; share_id; product_id','Private share'),('wishlist_feedback','share_id; user_id; reaction; comment','Feedback')],
 'screens':[('Personal Wishlist','Displays saved favourite products.','Allows removal from the list.','Keeps the list separate from cart.','Starts a product sharing action.'),('Friend Search and Requests','Lets a user find another account.','Creates a pending friend request.','Shows incoming and sent requests.','Supports acceptance or rejection.'),('Friends List','Shows accepted connections.','Provides an entry point for sharing.','Keeps pending requests distinct.','Represents the user social circle.'),('Shared Wishlist','Shows products sent by a friend.','Limits access to the intended recipient.','Preserves the sender and item context.','Allows feedback on the shared selection.'),('Share Feedback','Displays reactions and comments.','Lets a participant respond to the share.','Stores the response with the share record.','Supports collaborative style decisions.')]
},
{
 'n':'8.7','title':'Recommendations and Chatbot','scope':'This module helps shoppers discover complementary clothing and obtain quick fashion guidance. Outfit suggestions use predefined matching rules for eligible products; the chatbot uses deterministic intent and catalogue logic.',
 'fr':['Return a complementary outfit for an eligible base product.','Explain when an item is not eligible for outfit recommendations.','Answer supported product, stock, styling and shopping questions.','Offer links or guidance toward relevant catalogue and try-on pages.'],
 'nfr':['Do not present rule based results as learned predictions.','Use current catalogue values when answering product questions.','Return a controlled response for unsupported input or invalid product IDs.'],
 'stack':'Express recommendation engine and chatbot controller, React product and chatbot components, SQLite.',
 'flow':['The product page requests recommendations using the current product ID.','The recommendation engine checks eligibility and selects complementary product slots.','The chatbot receives a text message and matches supported intents.','The API returns suggestions or a bounded fallback response for display.'],
 'api':[('GET','/api/recommendations/product/:productId','Get outfit suggestions'),('POST','/api/chatbot/message','Get assistant response')],
 'db':[('products','id; category; color; price; stock','Current catalogue source'),('outfit_recommendations','base_product_id; recommended_product_id; slot_type','Stored recommendation links')],
 'screens':[('Complementary Product Suggestions','Shows products that complete a look.','Starts from an eligible base item.','Displays matching garments by slot.','Supports navigation to another product.'),('Complete Look Selection','Presents the coordinated set of items.','Shows how the items combine.','Supports an outfit purchase action.','Uses the rule based selection result.'),('Style Assistant Chat','Opens the assistant from the interface.','Accepts a shopping or styling question.','Returns a defined answer from application logic.','May guide the user to relevant features.')]
},
{
 'n':'8.8','title':'Virtual Try-On','scope':'Virtual try-on generates a preview of a supported garment on a user photograph. The Express API validates the request and contacts a separate local Python service that runs FASHN VTON 1.5.',
 'fr':['Report whether the local try-on service is ready.','Accept a supported photo and selected product.','Check garment eligibility and local garment image availability.','Send the person and garment images to the Python service.','Save and retrieve successful results for the authenticated user.'],
 'nfr':['Reject unsupported image types and missing uploads.','Limit upload size and protect result access by owner.','Report service unavailability without blocking ordinary shopping.'],
 'stack':'React virtual try-on page, Express and Multer, local Python Flask FASHN VTON 1.5, SQLite.',
 'flow':['The page checks VTON status and accepts an image upload.','Express validates the product, garment asset and person image.','The Python service generates an output image when ready.','Express saves the output path and result record, then returns it to the page.'],
 'api':[('GET','/api/vton/status','Check service readiness'),('POST','/api/vton/try-on','Generate try-on result'),('GET','/api/vton/history','List own results'),('GET','/api/vton/results/:id','Get one owned result')],
 'db':[('products','id; garment_image; vton_supported; vton_garment_category','Eligible garment metadata'),('try_on_results','user_id; product_id; generated_image_path; status','Generated result record')],
 'screens':[('Virtual Try-On Entry','Shows service readiness and instructions.','Lets the shopper select a supported product.','Accepts a person photograph.','Starts the local image generation request.'),('Photo and Garment Selection','Shows the chosen person image.','Uses the product garment image on disk.','Checks file type and required inputs.','Submits both images for inference.'),('Generated Try-On Result','Displays a successful generated image.','Associates the output with its product.','Makes the preview available to the same user.','Shows an error if generation fails.'),('Try-On History','Lists recent generated results.','Uses the signed-in user identifier.','Allows opening a saved result.','Does not expose another user result.')]
},
{
 'n':'8.9','title':'Admin Dashboard','scope':'The administrator module provides an account restricted view of catalogue and application activity. It offers summary metrics, users, purchases, virtual try-on activity and limited product updates.',
 'fr':['Require an ADMIN role before exposing routes.','Display metrics and recent activity.','List products, users and purchases for administration.','Update supported product fields through the admin route.','Show virtual try-on and email status information.'],
 'nfr':['Enforce the role on the backend even if a page route is hidden.','Avoid exposing password hashes in user summaries.','Validate editable product fields before writing them.'],
 'stack':'React admin pages, Express requireAdmin middleware, admin controller and SQLite.',
 'flow':['The admin page sends a JWT with dashboard requests.','The requireAdmin middleware verifies the account role.','Controller queries summarize current database records.','Product changes are validated and written through the admin endpoint.'],
 'api':[('GET','/api/admin/metrics','Read dashboard metrics'),('GET','/api/admin/products','List products'),('PATCH','/api/admin/products/:id','Update product'),('GET','/api/admin/users','List users'),('GET','/api/admin/purchases','List purchases'),('GET','/api/admin/vton-stats','Read try-on metrics')],
 'db':[('users','id; role; email','Role and user summary'),('products','id; price; stock; status','Catalogue editing'),('purchases','id; user_id; total_amount','Purchase summary'),('try_on_results','user_id; product_id; status','Try-on activity')],
 'screens':[('Admin Dashboard Overview','Displays restricted application summaries.','Shows current metrics from SQLite.','Provides navigation to management sections.','Requires an ADMIN account.'),('Admin Product Management','Lists catalogue products for review.','Shows editable product fields.','Submits supported changes to the API.','Reflects updated database values.'),('Admin Users and Purchases','Shows account and purchase summaries.','Supports operational visibility.','Keeps password hashes out of display.','Limits the view to administrators.'),('Admin Activity and VTON Metrics','Shows recent activity indicators.','Summarizes virtual try-on records.','Helps inspect service use and status.','Uses protected dashboard endpoints.')]
},
{
 'n':'8.10','title':'Theme and Language Preferences','scope':'This module lets a user select the application appearance, interface language and communication preferences. English, Tamil and Hindi translations are provided by the current frontend language context.',
 'fr':['Switch between available interface languages.','Remember the selected language across page reloads.','Change among supported appearance themes.','Save email, SMS and WhatsApp notification preferences.','Retrieve the signed-in user settings and delivery history.'],
 'nfr':['Use a supported language when a saved value is unavailable.','Keep account preferences associated with the correct user.','Avoid claiming external delivery from a preference switch alone.'],
 'stack':'React LanguageContext, ThemeContext and SettingsDrawer, Express settings routes, SQLite.',
 'flow':['The user opens settings and changes a language or appearance choice.','The frontend updates its context and local persistence for immediate display.','Account communication preferences are sent to the protected settings API.','The database stores the supported user preference fields.'],
 'api':[('GET','/api/settings','Read user settings'),('PUT','/api/settings','Save preferences'),('GET','/api/settings/notifications/history','Read delivery history')],
 'db':[('users','id; theme_preference; language_preference; email_notifications; sms_notifications; whatsapp_notifications','Saved user choices'),('notification_deliveries','user_id; channel; status','Communication history')],
 'screens':[('Settings Drawer','Groups appearance and communication choices.','Shows current saved values.','Allows changes from a single panel.','Provides feedback after a save.'),('Language Selector','Offers English, Tamil and Hindi.','Changes visible interface labels.','Remembers the choice across reloads.','Falls back when a translation key is absent.'),('Theme Selector','Offers the available appearance themes.','Applies the chosen style to the interface.','Keeps the selected preference.','Does not alter shopping records.'),('Communication Preferences','Shows channel switches for notifications.','Stores each enabled or disabled choice.','Controls attempted optional delivery.','Delivery success depends on provider configuration.')]
},
]

for m in modules:
    n=m['n']; sub(f"{n} {m['title']}")
    doc.add_heading(f'{n}.0 Purpose and Scope',3); body(m['scope'])
    doc.add_heading(f'{n}.1 Functional Requirements',3)
    for item in m['fr']: bullet(item)
    doc.add_heading(f'{n}.2 Non Functional Requirements',3)
    for item in m['nfr']: bullet(item)
    doc.add_heading(f'{n}.3 Tech Stack',3); body(m['stack'])
    doc.add_heading(f'{n}.4 Low Level Architecture',3)
    for item in m['flow']: bullet(item)
    doc.add_heading(f'{n}.5 API Design',3)
    table(['Method','Endpoint','Purpose'],m['api'],[2.5,8,5.5])
    doc.add_heading(f'{n}.6 DB Design',3)
    table(['Table','Key Fields','Use'],m['db'],[4,7,5])
    doc.add_heading(f'{n}.7 Screenshots',3)
    for i,(title,*lines) in enumerate(m['screens'],1):
        screenshot(f'{n}.{i}',title,'\n'.join(lines))

chapter(9,'CONCLUSION')
body('StyleHub brings the main operations of a men fashion shopping application into one system: protected accounts, product browsing, cart and purchase handling, wishlists, friend sharing, rule based suggestions, administration, and saved preferences. The SQLite schema and REST routes provide a coherent basis for these modules.')
body('Virtual try-on extends the project through a local third party model, while Google sign-in and communication channels depend on external configuration. The report diagrams and screenshot spaces support later submission updates without changing the underlying technical description.')

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('StyleHub Project Report')
all_paragraphs=list(doc.paragraphs)
for grid in doc.tables:
    for row in grid.rows:
        for cell in row.cells:
            all_paragraphs.extend(cell.paragraphs)
all_paragraphs.extend(sec.footer.paragraphs)
for paragraph in all_paragraphs:
    for run in paragraph.runs:
        run.font.name='Times New Roman'
        rfonts=run._element.get_or_add_rPr().rFonts
        if rfonts is not None:
            for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme'):
                rfonts.attrib.pop(qn(f'w:{attr}'),None)
            rfonts.set(qn('w:ascii'),'Times New Roman')
            rfonts.set(qn('w:hAnsi'),'Times New Roman')
            rfonts.set(qn('w:eastAsia'),'Times New Roman')
doc.save(OUT/'StyleHub_Project_Report.docx')
print('Created',*[str(x) for x in OUT.iterdir()],sep='\n')
