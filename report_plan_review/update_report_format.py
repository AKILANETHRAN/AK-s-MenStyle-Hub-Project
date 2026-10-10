from pathlib import Path
import re
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'report/StyleHub_Project_Report_Revised.docx'
OUTPUT = ROOT / 'report/StyleHub_Project_Report_Updated.docx'
d = Document(SOURCE)

STACKS = {
'8.2': [
('User interface and navigation', 'React and React Router', 'Provides registration, login, profile pages and protected navigation.'),
('Authentication API', 'Node.js and Express', 'Validates account requests and exposes authentication and profile endpoints.'),
('Session and role checks', 'jsonwebtoken (JWT)', 'Identifies the signed-in account and supports protected customer and administrator operations.'),
('Password protection', 'bcryptjs', 'Hashes passwords and verifies credentials without storing plaintext passwords.'),
('Account persistence', 'SQLite and better-sqlite3', 'Stores user identities, password hashes, usernames and account roles.')],
'8.3': [
('Catalogue interface', 'React', 'Renders product lists, details and recently viewed products.'),
('Product API', 'Node.js and Express', 'Serves catalogue data and records authenticated product views.'),
('Catalogue persistence', 'SQLite and better-sqlite3', 'Stores product information, prices, stock and recent views per account.')],
'8.4': [
('Cart state and checkout interface', 'React and CartContext', 'Keeps cart contents available across pages and displays checkout totals.'),
('Cart and checkout API', 'Node.js and Express', 'Handles item changes and validates checkout requests.'),
('Owner identification', 'JWT authentication', 'Associates cart operations with the signed-in user.'),
('Cart and purchase persistence', 'SQLite and better-sqlite3', 'Stores cart items and purchase records and supports transactional stock updates.')],
'8.5': [
('History and notification interface', 'React', 'Displays saved purchases and in-app notification status.'),
('Purchase and notification API', 'Node.js and Express', 'Processes purchases and retrieves account-specific history and messages.'),
('Transaction and delivery records', 'SQLite and better-sqlite3', 'Stores purchases, line items, notifications and channel delivery outcomes.'),
('Email communication', 'Nodemailer', 'Sends purchase email through a configured mail provider.'),
('SMS communication', 'Twilio', 'Sends purchase SMS when credentials and recipient settings are configured.'),
('WhatsApp communication', 'Meta WhatsApp Cloud API', 'Sends configured WhatsApp purchase messages and records delivery outcomes.')],
'8.6': [
('Social shopping interface', 'React', 'Displays wishlists, friend requests, shared products and feedback.'),
('Social API and access checks', 'Node.js, Express and JWT', 'Processes wishlist and friend actions with checks for owners and share participants.'),
('Relationship persistence', 'SQLite and better-sqlite3', 'Stores wishlists, friendships, shared items, reactions and comments.')],
'8.7': [
('Suggestions and chat interface', 'React', 'Shows complementary garments and supported assistant responses.'),
('Recommendation and chatbot API', 'Node.js and Express', 'Accepts product identifiers and chat messages and returns structured results.'),
('Matching and intent logic', 'JavaScript rule-based engine', 'Selects eligible outfit combinations and matches supported chat intents deterministically.'),
('Catalogue and recommendation data', 'SQLite and better-sqlite3', 'Supplies product attributes and stores recommendation records.')],
'8.8': [
('Image selection and results', 'React', 'Lets users select a photograph and garment and view generated results.'),
('Upload and request handling', 'Node.js, Express and Multer', 'Accepts multipart uploads and validates image and product inputs.'),
('Inference service interface', 'Python and Flask', 'Exposes the local service used by Express to request try-on generation.'),
('Virtual try-on generation', 'FASHN VTON 1.5', 'Generates a supported garment preview when model weights and hardware are available.'),
('Result persistence', 'SQLite and better-sqlite3', 'Stores result paths and ownership for account-specific try-on history.')],
'8.9': [
('Administration interface', 'React', 'Presents metrics, catalogue management and account activity.'),
('Administration API', 'Node.js and Express', 'Provides summary queries and validates supported product updates.'),
('Administrator authorization', 'JWT and requireAdmin middleware', 'Checks the authenticated account role before allowing management operations.'),
('Operational data', 'SQLite and better-sqlite3', 'Supplies product, user, purchase and try-on records for dashboard summaries.')],
'8.10': [
('Language selection', 'React LanguageContext', 'Applies English, Tamil or Hindi translations to the interface.'),
('Appearance selection', 'React ThemeContext', 'Applies the selected appearance across components.'),
('Settings interface', 'React SettingsDrawer', 'Groups language, theme and communication controls in one place.'),
('Browser persistence', 'Web Storage', 'Remembers supported interface choices across page reloads.'),
('Account settings API', 'Node.js, Express and JWT', 'Reads and saves communication preferences for the signed-in user.'),
('Account preference persistence', 'SQLite and better-sqlite3', 'Stores supported user preferences and channel delivery history.')]
}

def remove(p):
    p._element.getparent().remove(p._element)

def compact(p):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1

removed_folder = removed_low = joined = 0
paras = list(d.paragraphs)
for i,p in enumerate(paras):
    if not p.text.startswith('Fig '):
        continue
    following = paras[i+1]
    if following.text.startswith('This figure shows'):
        if p.text.startswith('Fig 8.1.'):
            remove(following); removed_folder += 1
        elif 'Low Level Architecture' in p.text:
            remove(following); removed_low += 1
        else:
            following.text = ' '.join(following.text.split())
            compact(following); joined += 1
    if p.text == 'Fig 8.1.2 Backend Folder':
        for extra in paras[i+2:i+5]:
            if extra.text.startswith('The backend contains'):
                remove(extra)

for p in list(d.paragraphs):
    m = re.fullmatch(r'(8\.\d+)\.3 Tech Stack', p.text)
    if not m: continue
    old = p._element.getnext()
    table = d.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.autofit = False
    widths = [Cm(3.8), Cm(4.7), Cm(7.5)]
    for col,w in zip(table.columns,widths): col.width = w
    for c,label,w in zip(table.rows[0].cells,['Concern','Choice','Why it fits this module'],widths):
        c.text = label; c.width = w
        for run in c.paragraphs[0].runs: run.bold = True
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7E6E6'); c._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in STACKS[m.group(1)]:
        cells = table.add_row().cells
        for c,value,w in zip(cells,row,widths): c.text=value; c.width=w
    for row_index,row in enumerate(table.rows):
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for c in row.cells:
            for para in c.paragraphs:
                compact(para)
                para.paragraph_format.keep_with_next = row_index < len(table.rows)-1
                for run in para.runs:
                    run.font.name = 'Times New Roman'; run.font.size = Pt(10); run.font.color.rgb = RGBColor(0,0,0)
    p._element.addnext(table._tbl)
    old.getparent().remove(old)

# Preserve the existing module-specific flow, format it as exactly four numbered points.
flows = 0
for i,p in enumerate(list(d.paragraphs)):
    if not (p.text.startswith('Fig ') and 'Low Level Architecture' in p.text): continue
    flows += 1
    flow_header = p.insert_paragraph_before('Request Flow')
    # Place the label immediately after the diagram caption.
    p._element.addnext(flow_header._element)
    compact(flow_header)
    flow_header.paragraph_format.keep_with_next = True
    for run in flow_header.runs: run.bold = True
    next_element = flow_header._element.getnext()
    for n in range(1,5):
        from docx.text.paragraph import Paragraph
        point = Paragraph(next_element,p._parent)
        assert point.text.startswith('•'), (p.text, point.text)
        point.text = str(n)+'. '+point.text.lstrip('• ').strip()
        compact(point)
        point.paragraph_format.keep_with_next = n<4
        next_element = next_element.getnext()

# Keep pictures and captions together, including intervening empty paragraphs.
all_paras = list(d.paragraphs)
for i,p in enumerate(all_paras):
    if not p.text.startswith('Fig '): continue
    if p.text.startswith('Fig 8.1.'):
        p.paragraph_format.keep_with_next = False
    else:
        p.paragraph_format.keep_with_next = True
    for preceding in reversed(all_paras[max(0,i-3):i]):
        if preceding.text.strip() and not preceding._p.xpath('.//w:drawing'): break
        preceding.paragraph_format.keep_with_next = True
        if preceding._p.xpath('.//w:drawing'): break

# The existing full-page class image needs room for its caption and paragraph.
class_image = d.inline_shapes[0]
class_image.width = int(class_image.width * 0.85)
class_image.height = int(class_image.height * 0.85)

d.save(OUTPUT)
assert removed_folder==12 and removed_low==9 and flows==9
assert sum(1 for t in d.tables if t.cell(0,0).text=='Concern')==9
print({'output':str(OUTPUT),'folder_descriptions_removed':removed_folder,'low_level_descriptions_removed':removed_low,'paragraph_descriptions':joined,'request_flows':flows,'module_tables':9})
