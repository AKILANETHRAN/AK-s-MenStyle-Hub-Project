from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches
from PIL import Image

ROOT = Path(r'D:\PROJECT\AK_MEN_STYLE_FULL_PROJECT')
SOURCE = Path(r'C:\Users\HARIRAAM\Documents\StyleHub_Project_Report_Final (2).docx')
OUTPUT = ROOT / 'StyleHub_Project_Report_Revised.docx'
ASSETS = ROOT / 'report_revision'

doc = Document(SOURCE)
original = list(doc.paragraphs)

# Only the screenshot explanations in Chapter 8 are expanded. The existing
# descriptions remain intact; each addition explains what its pictured UI
# demonstrates without implying unshown behaviour.
extensions = {
    128: ' The form keeps account access separate from the shopping pages until authentication succeeds. This screenshot also illustrates the entry point for users who already have an account and need to return to their saved activity.',
    130: ' Required identity and contact fields are grouped in one form so the customer can create a profile before using protected features. The page also offers a path back to sign-in for an existing account.',
    134: ' The account information is organized into editable sections rather than mixed with shopping controls. A completed username identifies the account in the application and enables access to account-specific features.',
    137: ' Its separate sign-in form makes the administrative entry point clear. The screen itself is only the first step; the backend role check determines whether the signed-in account may use management functions.',
    156: ' The featured-product area provides visual cards that connect the introductory page to individual product records. A shopper can use the navigation and featured selections to continue toward the full catalogue.',
    160: ' Filters and product cards help narrow the visible selection before a product is opened. Each card gives enough information for an initial comparison, while the detail page holds the size and purchase controls.',
    163: ' The larger garment image and product information support a closer inspection before a decision is made. The available actions connect the same catalogue record to the customer’s shopping and styling workflows.',
    182: ' This view shows the purchase options on the product page and the size selection needed for a clothing item. The user can confirm the selection before continuing to the cart and later checkout.',
    185: ' Item rows keep the chosen product and quantity visible while the order is being prepared. The shopper can revise the cart rather than restarting the product search when a size or quantity needs adjustment.',
    188: ' The reward area makes the applicable discount visible next to the order calculation. This helps the shopper understand the amount payable before moving from cart review to the secure checkout form.',
    191: ' Delivery information and the payment summary are presented together for a final review. The screenshot documents the last customer-facing step before the backend creates the purchase and its item records.',
    210: ' Each purchase entry connects an order identifier with its amount and saved delivery information. This gives the customer a record of completed transactions without exposing the purchases of another signed-in user.',
    213: ' The panel keeps short status messages accessible from the main interface. It distinguishes the in-app record from an external email, SMS or WhatsApp attempt, which may have a separate delivery outcome.',
    217: ' The email presents the purchase as a readable receipt with the confirmed products and recipient details. It is an example of the customer communication produced after the purchase record is created.',
    220: ' The screenshot documents a provider test message and a separate StyleHub product alert, not a purchase receipt. It demonstrates the WhatsApp channel and the type of information a configured integration can deliver.',
    239: ' Product cards let the user compare saved items and reopen a product from the personal list. Saving an item here does not place it in the cart or create a purchase; those actions remain separate.',
    242: ' Search results show matching accounts and the available request action. A sent request remains distinct from an accepted friendship, so sharing controls are based on the relationship state.',
    245: ' Accepted connections are shown together in the social area. This gives the user a clear destination for selecting a friend before sharing one or more saved fashion items.',
    248: ' The dialog identifies the chosen wishlist products and the intended recipient before the share is sent. Keeping that choice explicit helps prevent a personal wishlist from being exposed to the wrong account.',
    251: ' The screenshot also shows how the receiver can evaluate the shared selection in context. Reactions and written feedback support a style discussion without altering the sender’s original wishlist items.',
    269: ' The conversation panel presents suggested prompts and product-oriented responses in one place. Its answers are based on application rules and catalogue information rather than on a trained conversational model.',
    272: ' The recommendation area displays complementary products beside the selected base item. Images, prices and product names let the shopper inspect the suggested outfit before choosing any additional action.',
    313: ' The two panels keep the uploaded person photograph separate from the garment reference used for inference. This lets the user verify both inputs before asking the local try-on service to generate a preview.',
    316: ' The result remains connected to the photograph and selected garment shown in the preceding step. These controls make it possible to keep the output or repeat the process with a different input.',
    320: ' This selection view records the second shirt test before generation. Showing the original photo next to the isolated denim shirt makes the intended clothing change clear without changing the person input.',
    323: ' The shirt colour and shape can be compared with the garment-only image from the selection screen. The available actions document how the generated preview can be retained or produced again.',
    327: ' This screenshot demonstrates the lower-body input path: the uploaded person is shown beside a trouser-only product image. The product details and service-status badges help confirm that the correct garment and engine are selected.',
    330: ' The result can be compared with the input photograph to see the clothing change while the person and surrounding scene remain recognizable. The controls support saving this preview or beginning another try-on.',
    349: ' Product rows bring catalogue information and edit actions into one administrative view. The screenshot illustrates that management is performed on stored products rather than through the customer-facing detail page.',
    352: ' Separate lists allow the administrator to inspect account records and completed purchases. These views support oversight of the same users and transactions that appear in the application’s customer workflows.',
    355: ' The counts provide a compact status summary instead of showing individual generated images. This helps an administrator see recent usage and outcomes while keeping the detailed try-on flow on the customer side.',
    375: ' The three choices are visible together so a user can select a preferred interface language. The active choice is indicated in the control, making the current language easy to check before leaving Settings.',
    378: ' The available colour themes are presented as separate selectable options. This screen demonstrates the appearance preference control, while the selected theme is also applied across the shopping interface.',
    381: ' Each communication channel has its own control, allowing the user to choose which purchase-related notifications are enabled. The preference is saved to the account, but actual delivery still depends on provider configuration.'
}

for index, addition in extensions.items():
    paragraph = original[index]
    run = paragraph.add_run(addition)
    prior = next((r for r in reversed(paragraph.runs[:-1]) if r.text.strip()), None)
    if prior is not None and prior._element.rPr is not None:
        run._element.insert(0, deepcopy(prior._element.rPr))

original[220].text = (
    'The screenshot shows a Meta sample message confirming the test API connection, followed by a StyleHub '
    'new-product alert with its price. This is a test-recipient notification rather than a purchase receipt or '
    'customer broadcast. It illustrates the WhatsApp channel and the kind of product information a configured '
    'integration can deliver.'
)

# The original feedback description begins with an accidental backtick.
for run in original[251].runs:
    if run.text == '`':
        run.text = ''

def insert_after(reference, text, style=None):
    element = OxmlElement('w:p')
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph

insert_after(original[108],
    'The project folder view shows the top-level separation between the frontend and backend applications. Keeping these areas distinct makes it easier to find the browser interface, server APIs and related project resources. The structure shown here is the starting point for the more detailed module folders below.')
insert_after(original[109],
    'The backend folder groups its configuration, routes, controllers and services around API responsibilities, while the frontend folder holds the React application and its assets. These side-by-side views explain where the report’s browser pages and server operations are implemented. They also make the boundary between the two parts of the project visible.')

# The supplied Chapter 8 skips from 8.10 to 8.20 and from 8.44 to 8.46.
# Renumber caption text only, leaving images and other chapter references alone.
def fix_caption(match):
    old = int(match.group(1))
    if 20 <= old <= 44:
        return f'Fig 8.{old - 9}'
    if 46 <= old <= 48:
        return f'Fig 8.{old - 10}'
    return match.group(0)

for paragraph in original[103:382]:
    for run in paragraph.runs:
        if 'Fig 8.' in run.text:
            run.text = re.sub(r'Fig 8\.(\d+)', fix_caption, run.text)

# The VTON comparison table uses static figure references, not caption fields.
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text in {'Fig 8.37', 'Fig 8.39', 'Fig 8.41'}:
                        run.text = {'Fig 8.37': 'Fig 8.28',
                                    'Fig 8.39': 'Fig 8.30',
                                    'Fig 8.41': 'Fig 8.32'}[run.text]

# The source embeds the second shirt-selection image at the end of the prior
# screenshot's description. Move that drawing into the existing blank image
# paragraph, so Word can keep it with its own caption and explanation.
for run in list(original[310].runs):
    if run._element.xpath('.//w:drawing'):
        original[310]._p.remove(run._element)
        original[311]._p.append(run._element)
original[310].text = (
    'The fitting room is shown before a photograph is selected. Its progress strip marks the photo upload as '
    'the active first step, while the garment-only shirt image identifies the product chosen for try-on. The '
    'upload panel accepts JPG, JPEG, PNG or WEBP images up to 15 MB. Generation remains unavailable until the '
    'customer adds a suitable photograph.'
)
original[311].paragraph_format.keep_with_next = True
original[312].paragraph_format.keep_with_next = True
original[129].paragraph_format.keep_with_next = True

# Preserve the two-paragraph Conclusion and move the proposed work into 9.1.
original[384].text = (
    'StyleHub combines protected accounts, product discovery, carts and purchases, wishlists, friend sharing, '
    'rule-based styling assistance, administration and saved preferences in one men’s fashion shopping platform. '
    'A shared product catalogue and authenticated account connect these modules, while React pages, Express APIs '
    'and SQLite records keep customer actions and data organized. Password hashes, role checks and account-scoped '
    'queries support the security of shopping and social features.'
)
original[385].text = (
    'Virtual try-on adds an optional preview through a separate local FASHN VTON service, and external sign-in or '
    'messaging works only when its provider configuration is available. The current recommender and chatbot use '
    'defined rules rather than learned predictions, and SQLite suits the documented single-server project more '
    'than heavy concurrent deployment. Within these limits, the implemented modules support a complete path from '
    'finding a garment to considering it with friends, trying it on and purchasing it.'
)

future = (
    'Future work could add an online payment gateway and order tracking, migrate the primary database to a '
    'server-based system for larger deployments, and improve outfit recommendations using responsibly collected '
    'browsing and purchase history. A hosted GPU service could make virtual try-on available beyond a local '
    'machine. A broader automated test suite, additional interface languages and more robust provider monitoring '
    'would further improve reliability and accessibility. These are proposed enhancements, not features claimed '
    'as part of the current implementation.'
)
for paragraph in original[386:391]:
    paragraph._element.getparent().remove(paragraph._element)
heading = insert_after(original[385], '9.1 Future Enhancement', 'Heading 2')
insert_after(heading, future)

# Replace only the five dark VTON screenshots. ImageGen-derived assets have
# been visually checked against their source for people, garments and labels.
for rid in range(42, 47):
    part = doc.part.related_parts[f'rId{rid}']
    part._blob = (ASSETS / f'rId{rid}-light.png').read_bytes()

# Bring individual screenshots into a consistent print-size envelope, without
# cropping or stretching. Paired screenshots retain their side-by-side layout.
paired = {14, 15, 20, 21, 33, 34}
sizes = []
for shape in doc.inline_shapes:
    blips = shape._inline.xpath('.//a:blip')
    if not blips:
        continue
    rid = blips[0].get(qn('r:embed'))
    if not rid or not rid.startswith('rId'):
        continue
    number = int(rid[3:])
    if not 13 <= number <= 52 or number in paired:
        continue
    blob = doc.part.related_parts[rid].blob
    with Image.open(BytesIO(blob)) as image:
        ratio = image.width / image.height
    if ratio > 3:
        width = 5.5
        height = width / ratio
    elif ratio >= 1.1:
        width = 5.1
        height = width / ratio
        if height > 3.5:
            height = 3.5
            width = height * ratio
    else:
        height = 3.5
        width = height * ratio
        if width > 5.1:
            width = 5.1
            height = width / ratio
    if number == 16:  # login + registration should fit with both explanations
        height, width = 2.45, 2.45 * ratio
    elif number == 17:
        height, width = 3.5, 3.5 * ratio
    shape.width = Inches(width)
    shape.height = Inches(height)
    sizes.append((rid, round(width, 2), round(height, 2)))

captions = []
for paragraph in doc.paragraphs:
    captions.extend(int(n) for n in re.findall(r'Fig 8\.(\d+)', paragraph.text))
assert captions == list(range(1, 39)), captions
assert len([p for p in doc.paragraphs if p.text == 'CONCLUSION']) == 1
assert len([p for p in doc.paragraphs if p.text == '9.1 Future Enhancement']) == 1

# Update the cached Chapter 9 TOC page to match the final Word render.
toc_page = doc.tables[0].cell(9, 2)
for text_node in toc_page._tc.xpath('.//w:t'):
    if text_node.text == '57':
        text_node.text = '54'

doc.save(OUTPUT)
print(f'Saved {OUTPUT}')
print('Chapter 8 figures:', captions)
print('Resized screenshots:', sizes)
